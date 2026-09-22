from __future__ import annotations

"""
SCRIPT 07.3 v1.0.0 — REGION VS COUNTRY

Purpose
-------
Answer the research question behind the dissertation's continuation: does it
pay off to compare regions rather than countries? A region can have
particularities that other regions of the same country do not share, and can
resemble regions of other countries more closely.

This script performs NO re-estimation of the frozen analytical model. It only
reads already-validated outputs of Script 07.2 (outputs/07.2_powerbi_model.xlsx)
and computes three new analyses on top of them:

1) Variance decomposition — how much of each of 14 region-level variables is
   explained by country membership (eta-squared, bias-corrected epsilon-
   squared, permutation significance).
2) Region-to-region neighbors across borders — for each region, whether its
   most similar regions (by trajectory distance) are in the same country or
   abroad, against a country-size baseline.
3) Country-as-unit counterfactual — what a region would "inherit" if its
   country's majority profile were used instead of its own, and whether
   countries contain regions moving in opposite directions at once.

Official P1 state expected (see Script 07.2 for the full frozen state)
-----------------------------------------------------------------------
- Analytical period: 2015-2019.
- Final analytical sample: 512 regions in 28 countries.
- Hard clusters: C1=100, C2=412.

Interpretation guardrail
------------------------
Findings here are descriptive, not causal. A high country share of variation
in the assigned profile does NOT mean "regions are better than countries," and
a region's most similar region being abroad does NOT mean it "belongs" to
that other country's profile. See RVC_Metadata / results/07.3_results_report.md
for the full set of interpretation guardrails that must be respected in any
report or dashboard text built on these results.
"""

from pathlib import Path
import hashlib
import json
import time
from typing import Any

import numpy as np
import pandas as pd


# =============================================================================
# CONFIGURATION
# =============================================================================

ROOT = Path(__file__).resolve().parents[1]
OUTPUTS_DIR = ROOT / "outputs"
RESULTS_DIR = ROOT / "results"

SOURCE_WORKBOOK = OUTPUTS_DIR / "07.2_powerbi_model.xlsx"
WORKBOOK_OUTPUT = OUTPUTS_DIR / "07.3_region_vs_country.xlsx"
REPORT_OUTPUT = RESULTS_DIR / "07.3_results_report.md"
FINAL_STATUS_OUTPUT = ROOT / "data" / "interim" / "script07" / "07.3_final_status.json"

SCRIPT_VERSION = "1.0.0"

SEED = 20260922
B_PERM = 9999
K_LIST = (1, 5, 10)
EXPECTED_REGIONS = 512
EXPECTED_COUNTRIES = 28
EXPECTED_YEARS = [2015, 2016, 2017, 2018, 2019]
DISTANCE_FEATURES = (
    "bod_demographic_productive_potential",
    "bod_socioeconomic_deprivation",
    "ntl_per_capita_norm",
)
NUMERIC_TOL_DISTANCE = 1e-12
QA_TOLERANCE = 1e-3

REQUIRED_SHEETS = (
    "Fact_Region",
    "Region_Year_Feature",
    "Region_Year_Indicator",
    "Region_Country_Ref",
    "Country_Ref_Region",
    "Global_Local_Country",
    "Dim_Country",
)

# -----------------------------------------------------------------------------
# Analysis 1 — variables under decomposition, in A.2 order
# -----------------------------------------------------------------------------

VARIANCE_VARIABLES = [
    {"variable": "hard_cluster_is_C1", "variable_label": "Assigned profile (1 = C1)",
     "variable_group": "profile_level", "source": "fact_region_flag", "column": "hard_cluster"},
    {"variable": "u_cluster_1", "variable_label": "Membership in C1",
     "variable_group": "profile_level", "source": "fact_region", "column": "u_cluster_1"},
    {"variable": "membership_margin", "variable_label": "Membership margin",
     "variable_group": "profile_level", "source": "fact_region", "column": "membership_margin"},
    {"variable": "convergence_index", "variable_label": "Convergence index",
     "variable_group": "movement", "source": "fact_region", "column": "convergence_index"},
    {"variable": "relative_gap_2015", "variable_label": "Relative position 2015",
     "variable_group": "movement", "source": "fact_region", "column": "relative_gap_2015"},
    {"variable": "relative_gap_2019", "variable_label": "Relative position 2019",
     "variable_group": "movement", "source": "fact_region", "column": "relative_gap_2019"},
    {"variable": "bod_socioeconomic_deprivation", "variable_label": "Socioeconomic deprivation composite",
     "variable_group": "dimension", "source": "region_year_feature", "column": "bod_socioeconomic_deprivation"},
    {"variable": "bod_demographic_productive_potential", "variable_label": "Demographic productive potential",
     "variable_group": "dimension", "source": "region_year_feature", "column": "bod_demographic_productive_potential"},
    {"variable": "ntl_per_capita_norm", "variable_label": "Night-time lights per capita",
     "variable_group": "indicator", "source": "region_year_feature", "column": "ntl_per_capita_norm"},
    {"variable": "poor420", "variable_label": "Poverty (poor420)",
     "variable_group": "indicator", "source": "region_year_indicator", "column": "poor420"},
    {"variable": "gini", "variable_label": "Gini coefficient",
     "variable_group": "indicator", "source": "region_year_indicator", "column": "gini"},
    {"variable": "prosgap2021", "variable_label": "Prosperity gap (prosgap2021)",
     "variable_group": "indicator", "source": "region_year_indicator", "column": "prosgap2021"},
    {"variable": "share_15_64", "variable_label": "Population aged 15-64",
     "variable_group": "indicator", "source": "region_year_indicator", "column": "share_15_64"},
    {"variable": "share_65_plus", "variable_label": "Population aged 65+",
     "variable_group": "indicator", "source": "region_year_indicator", "column": "share_65_plus"},
]

# -----------------------------------------------------------------------------
# A.6 QA — expected values (Section A.6 of the specification)
# -----------------------------------------------------------------------------

EXPECTED_VARIANCE = {
    "hard_cluster_is_C1": (0.697, 0.680),
    "u_cluster_1": (0.680, 0.663),
    "membership_margin": (0.432, 0.400),
    "convergence_index": (0.121, 0.072),
    "relative_gap_2015": (0.300, 0.261),
    "relative_gap_2019": (0.268, 0.227),
    "bod_socioeconomic_deprivation": (0.774, 0.761),
    "bod_demographic_productive_potential": (0.567, 0.543),
    "ntl_per_capita_norm": (0.289, 0.249),
    "poor420": (0.769, 0.756),
    "gini": (0.835, 0.826),
    "prosgap2021": (0.792, 0.781),
    "share_15_64": (0.509, 0.482),
    "share_65_plus": (0.875, 0.868),
}
EXPECTED_NULL_ETA2 = 0.0528
EXPECTED_NN1_FOREIGN_COUNT = 179
EXPECTED_NN1_FOREIGN_SHARE = 0.350
EXPECTED_MEAN_FOREIGN_K5 = 0.426
EXPECTED_MAJORITY_FOREIGN_K5 = 214
EXPECTED_MEAN_FOREIGN_K10 = 0.489
EXPECTED_MAJORITY_FOREIGN_K10 = 222
EXPECTED_MEAN_BASELINE_FOREIGN = 0.922
EXPECTED_PERM_MEAN_APPROX = 0.922
EXPECTED_PERM_P_MAX = 0.001
EXPECTED_DIFFERS_MAJORITY = 37
EXPECTED_DIFFERS_MAJORITY_PCT = 0.072
EXPECTED_BOTH_DIRECTIONS_COUNTRIES = 24
EXPECTED_OWN_COUNTRY_NEAREST = 188


# =============================================================================
# BASIC HELPERS (style mirrors Script 07.1 / 07.2)
# =============================================================================

def print_header(title: str) -> None:
    print("\n" + "=" * 100)
    print(title)
    print("=" * 100)


def ensure_directories() -> None:
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    FINAL_STATUS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)


def relative_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.name


def file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def qa_add(rows: list[dict[str, Any]], check: str, condition: bool,
           observed: Any, expected: Any, detail: str = "") -> None:
    rows.append({
        "check": check,
        "status": "PASS" if condition else "FAIL",
        "observed": observed,
        "expected": expected,
        "detail": detail,
    })


def enforce_qa(rows: list[dict[str, Any]], context: str) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    failed = frame[frame["status"].eq("FAIL")]
    if not failed.empty:
        message = failed[["check", "observed", "expected", "detail"]].to_string(index=False)
        raise RuntimeError(
            f"{context} failed — STOP POINT reached; Part B must NOT proceed:\n{message}"
        )
    return frame


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8-sig")


def autofit_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        sample_rows = min(ws.max_row, 500)
        for column_cells in ws.iter_cols(min_row=1, max_row=sample_rows):
            letter = column_cells[0].column_letter
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 48)
    wb.save(path)


# =============================================================================
# 1/7 — INPUT CONTRACT
# =============================================================================

def build_input_manifest() -> pd.DataFrame:
    if not SOURCE_WORKBOOK.exists():
        raise FileNotFoundError(
            f"Script 07.3 cannot start; required input is missing: {relative_path(SOURCE_WORKBOOK)}"
        )
    rows = [{
        "logical_name": "07.2_powerbi_model",
        "relative_path": relative_path(SOURCE_WORKBOOK),
        "size_bytes": SOURCE_WORKBOOK.stat().st_size,
        "sha256": file_sha256(SOURCE_WORKBOOK),
    }]
    return pd.DataFrame(rows)


def load_inputs() -> dict[str, pd.DataFrame]:
    xl = pd.ExcelFile(SOURCE_WORKBOOK)
    missing = [s for s in REQUIRED_SHEETS if s not in xl.sheet_names]
    if missing:
        raise RuntimeError(f"07.2 workbook is missing required sheet(s): {missing}")
    return {name: xl.parse(name) for name in REQUIRED_SHEETS}


def validate_input_contract(t: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    fact_region = t["Fact_Region"]
    qa_add(rows, "fact_region_rows", len(fact_region) == EXPECTED_REGIONS,
           len(fact_region), EXPECTED_REGIONS)
    qa_add(rows, "fact_region_countries", fact_region["code"].nunique() == EXPECTED_COUNTRIES,
           fact_region["code"].nunique(), EXPECTED_COUNTRIES)
    qa_add(rows, "fact_region_unique_ids", fact_region["region_id"].is_unique,
           fact_region["region_id"].is_unique, True)

    ryf = t["Region_Year_Feature"]
    qa_add(rows, "region_year_feature_rows",
           len(ryf) == EXPECTED_REGIONS * len(EXPECTED_YEARS) * len(DISTANCE_FEATURES),
           len(ryf), EXPECTED_REGIONS * len(EXPECTED_YEARS) * len(DISTANCE_FEATURES))

    ryi = t["Region_Year_Indicator"]
    indicator_vars = {"poor420", "gini", "prosgap2021", "share_15_64", "share_65_plus"}
    observed_vars = set(ryi["variable"].unique())
    qa_add(rows, "region_year_indicator_variables", indicator_vars.issubset(observed_vars),
           sorted(observed_vars), sorted(indicator_vars))

    rcr = t["Region_Country_Ref"]
    qa_add(rows, "region_country_ref_rows",
           len(rcr) == EXPECTED_REGIONS * EXPECTED_COUNTRIES,
           len(rcr), EXPECTED_REGIONS * EXPECTED_COUNTRIES)

    ccr = t["Country_Ref_Region"]
    qa_add(rows, "country_ref_region_rows", len(ccr) == EXPECTED_REGIONS,
           len(ccr), EXPECTED_REGIONS)

    return enforce_qa(rows, "Script 07.3 input contract")


# =============================================================================
# 2/7 — ANALYSIS 1: VARIANCE DECOMPOSITION (how much does the country explain)
# =============================================================================

def build_variable_matrix(t: dict[str, pd.DataFrame], region_order: pd.Series) -> pd.DataFrame:
    """Return a (512 x 14) frame indexed by region_id, columns = VARIANCE_VARIABLES order."""
    fact_region = t["Fact_Region"].set_index("region_id")
    ryf = t["Region_Year_Feature"]
    ryi = t["Region_Year_Indicator"]

    feature_means = (
        ryf[ryf["feature"].isin(DISTANCE_FEATURES)]
        .pivot_table(index="region_id", columns="feature", values="feature_value", aggfunc="mean")
    )
    indicator_vars = ("poor420", "gini", "prosgap2021", "share_15_64", "share_65_plus")
    indicator_means = (
        ryi[ryi["variable"].isin(indicator_vars)]
        .pivot_table(index="region_id", columns="variable", values="value_normalized", aggfunc="mean")
    )

    columns: dict[str, pd.Series] = {}
    for spec in VARIANCE_VARIABLES:
        name = spec["variable"]
        source = spec["source"]
        column = spec["column"]
        if source == "fact_region_flag":
            series = (fact_region[column] == 1).astype(float)
        elif source == "fact_region":
            series = fact_region[column].astype(float)
        elif source == "region_year_feature":
            series = feature_means[column]
        elif source == "region_year_indicator":
            series = indicator_means[column]
        else:
            raise RuntimeError(f"Unknown source: {source}")
        columns[name] = series.reindex(region_order)

    matrix = pd.DataFrame(columns, index=region_order)
    if matrix.isna().any().any():
        bad = matrix.columns[matrix.isna().any()].tolist()
        raise RuntimeError(f"Missing values while building variance matrix for: {bad}")
    return matrix


def eta2_epsilon2_from_sums(sums: np.ndarray, n_g: np.ndarray, totals: np.ndarray,
                            sst: np.ndarray, n: int, k: int) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    ssb = (sums ** 2 / n_g[:, None]).sum(axis=0) - (totals ** 2) / n
    eta2 = ssb / sst
    msw = (sst - ssb) / (n - k)
    epsilon2 = (ssb - (k - 1) * msw) / sst
    return eta2, epsilon2, ssb


def run_variance_decomposition(matrix: pd.DataFrame, code: pd.Series) -> pd.DataFrame:
    n = len(matrix)
    countries, group_idx = np.unique(code.to_numpy(), return_inverse=True)
    k = len(countries)
    n_g = np.bincount(group_idx, minlength=k).astype(float)

    X = matrix.to_numpy(dtype=float)
    grand_means = X.mean(axis=0)
    sst = ((X - grand_means) ** 2).sum(axis=0)
    totals = X.sum(axis=0)

    order = np.argsort(group_idx, kind="stable")
    sorted_group_idx = group_idx[order]
    boundaries = np.concatenate(([0], np.cumsum(n_g)[:-1])).astype(int)

    obs_sums = np.add.reduceat(X[order], boundaries, axis=0)
    # Sanity: obs_sums row g should equal the sum of X restricted to group g.
    eta2_obs, epsilon2_obs, _ = eta2_epsilon2_from_sums(obs_sums, n_g, totals, sst, n, k)

    rng = np.random.default_rng(SEED)
    n_vars = X.shape[1]
    exceed_counts = np.zeros(n_vars, dtype=np.int64)
    for _ in range(B_PERM):
        perm = rng.permutation(n)
        Xp = X[perm][order]
        sums_p = np.add.reduceat(Xp, boundaries, axis=0)
        eta2_p, _, _ = eta2_epsilon2_from_sums(sums_p, n_g, totals, sst, n, k)
        exceed_counts += (eta2_p >= eta2_obs).astype(np.int64)

    perm_p = (1.0 + exceed_counts) / (B_PERM + 1)
    null_expected_eta2 = (k - 1) / (n - 1)
    within_share = 1.0 - eta2_obs

    out = pd.DataFrame({
        "variable": [s["variable"] for s in VARIANCE_VARIABLES],
        "variable_label": [s["variable_label"] for s in VARIANCE_VARIABLES],
        "variable_group": [s["variable_group"] for s in VARIANCE_VARIABLES],
        "n_regions": n,
        "n_countries": k,
        "eta2": eta2_obs,
        "epsilon2": epsilon2_obs,
        "within_share": within_share,
        "perm_p": perm_p,
        "perm_B": B_PERM,
        "null_expected_eta2": null_expected_eta2,
    })
    return out


# =============================================================================
# 3/7 — ANALYSIS 2: NEIGHBORS ACROSS BORDERS
# =============================================================================

def build_feature_year_matrix(t: dict[str, pd.DataFrame], region_order: pd.Series) -> np.ndarray:
    """Return (512, 5, 3) array: region x year x feature (fixed year/feature order)."""
    ryf = t["Region_Year_Feature"]
    pivot = ryf.pivot_table(
        index="region_id", columns=["year", "feature"], values="feature_value", aggfunc="mean"
    )
    cols = pd.MultiIndex.from_product([EXPECTED_YEARS, DISTANCE_FEATURES])
    pivot = pivot.reindex(index=region_order, columns=cols)
    if pivot.isna().any().any():
        raise RuntimeError("Missing region/year/feature combinations in Region_Year_Feature.")
    arr = pivot.to_numpy(dtype=float).reshape(len(region_order), len(EXPECTED_YEARS), len(DISTANCE_FEATURES))
    return arr


def validate_country_reference_distance(
    F: np.ndarray, region_order: pd.Series, code: pd.Series, t: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    """Reproduce Region_Country_Ref.mean_annual_distance_2015_2019 (14,336 rows) exactly."""
    n = len(region_order)
    countries = np.sort(code.unique())
    region_pos = {rid: i for i, rid in enumerate(region_order)}
    code_arr = code.to_numpy()

    # Full-country mean per year/feature, and per-region leave-one-out own-country mean.
    country_of = {c: np.where(code_arr == c)[0] for c in countries}

    n_years, n_feat = len(EXPECTED_YEARS), len(DISTANCE_FEATURES)
    recomputed = np.zeros((n, len(countries)))
    recomputed_2015 = np.zeros((n, len(countries)))
    recomputed_2019 = np.zeros((n, len(countries)))

    for ci, c in enumerate(countries):
        idx_c = country_of[c]
        sum_c = F[idx_c].sum(axis=0)  # (n_years, n_feat)
        n_c = len(idx_c)
        full_mean = sum_c / n_c  # (n_years, n_feat), reference for non-home regions

        # Distance from every region to the full-country mean.
        d_full = ((F - full_mean[None, :, :]) ** 2).mean(axis=2)  # (n, n_years)

        # Leave-one-out mean, applied only to regions that belong to country c.
        if n_c > 1:
            loo_mean = (sum_c[None, :, :] - F[idx_c]) / (n_c - 1)  # (n_c, n_years, n_feat)
            d_loo = ((F[idx_c] - loo_mean) ** 2).mean(axis=2)  # (n_c, n_years)
            d_full[idx_c] = d_loo
        # (n_c == 1 never occurs in this sample; no self-only country.)

        recomputed[:, ci] = d_full.mean(axis=1)
        recomputed_2015[:, ci] = d_full[:, EXPECTED_YEARS.index(2015)]
        recomputed_2019[:, ci] = d_full[:, EXPECTED_YEARS.index(2019)]

    long = pd.DataFrame(recomputed, index=region_order, columns=countries)
    long = long.stack().rename("recomputed_mean_annual_distance").reset_index()
    long.columns = ["region_id", "reference_country_code", "recomputed_mean_annual_distance"]

    rcr = t["Region_Country_Ref"][
        ["region_id", "reference_country_code", "mean_annual_distance_2015_2019"]
    ]
    merged = rcr.merge(long, on=["region_id", "reference_country_code"], how="left", validate="one_to_one")
    merged["abs_diff"] = (
        merged["mean_annual_distance_2015_2019"] - merged["recomputed_mean_annual_distance"]
    ).abs()
    return merged


def run_neighbor_analysis(
    F: np.ndarray, region_order: pd.Series, code: pd.Series, t: dict[str, pd.DataFrame],
) -> tuple[pd.DataFrame, dict[str, Any]]:
    n = len(region_order)
    Fflat = F.reshape(n, -1)  # (512, 15) — 5 years x 3 features

    diff = Fflat[:, None, :] - Fflat[None, :, :]
    D = (diff ** 2).mean(axis=2)
    np.fill_diagonal(D, np.inf)

    order_nn = np.argsort(D, axis=1, kind="stable")  # stable -> ties broken by region_id order
    code_arr = code.to_numpy()
    region_arr = region_order.to_numpy()

    countries, n_per_country_arr = np.unique(code_arr, return_counts=True)
    n_c_lookup = dict(zip(countries, n_per_country_arr))
    n_c = np.array([n_c_lookup[c] for c in code_arr])
    baseline_foreign_share = 1.0 - (n_c - 1) / (n - 1)

    foreign_shares = {}
    for k in K_LIST:
        nn_k = order_nn[:, :k]
        neighbor_codes = code_arr[nn_k]
        is_foreign = neighbor_codes != code_arr[:, None]
        foreign_shares[k] = is_foreign.mean(axis=1)

    nn1_idx = order_nn[:, 0]
    nn1_region_id = region_arr[nn1_idx]
    nn1_country_code = code_arr[nn1_idx]
    nn1_distance = D[np.arange(n), nn1_idx]
    nn1_is_foreign = nn1_country_code != code_arr

    nearest_same_country_region_id = np.full(n, "", dtype=object)
    nearest_same_country_distance = np.full(n, np.nan)
    nearest_foreign_region_id = np.full(n, "", dtype=object)
    nearest_foreign_country_code = np.full(n, "", dtype=object)
    nearest_foreign_distance = np.full(n, np.nan)

    for i in range(n):
        row_order = order_nn[i]
        row_codes = code_arr[row_order]
        same_mask = row_codes == code_arr[i]
        if same_mask.any():
            j = row_order[same_mask][0]
            nearest_same_country_region_id[i] = region_arr[j]
            nearest_same_country_distance[i] = D[i, j]
        foreign_mask = ~same_mask
        if foreign_mask.any():
            j = row_order[foreign_mask][0]
            nearest_foreign_region_id[i] = region_arr[j]
            nearest_foreign_country_code[i] = code_arr[j]
            nearest_foreign_distance[i] = D[i, j]

    foreign_closer_than_same_country = nearest_foreign_distance < nearest_same_country_distance
    k5_exceeds_same_country_pool = (n_c - 1) < 5
    k10_exceeds_same_country_pool = (n_c - 1) < 10

    region_out = pd.DataFrame({
        "region_id": region_arr,
        "code": code_arr,
        "nn1_region_id": nn1_region_id,
        "nn1_country_code": nn1_country_code,
        "nn1_distance": nn1_distance,
        "nn1_is_foreign": nn1_is_foreign,
        "foreign_share_k1": foreign_shares[1],
        "foreign_share_k5": foreign_shares[5],
        "foreign_share_k10": foreign_shares[10],
        "baseline_foreign_share": baseline_foreign_share,
        "k5_exceeds_same_country_pool": k5_exceeds_same_country_pool,
        "k10_exceeds_same_country_pool": k10_exceeds_same_country_pool,
        "nearest_same_country_region_id": nearest_same_country_region_id,
        "nearest_same_country_distance": nearest_same_country_distance,
        "nearest_foreign_region_id": nearest_foreign_region_id,
        "nearest_foreign_country_code": nearest_foreign_country_code,
        "nearest_foreign_distance": nearest_foreign_distance,
        "foreign_closer_than_same_country": foreign_closer_than_same_country,
    })

    # Permutation test on the mean of foreign_share_k5 (Section A.3.6).
    nn5_idx = order_nn[:, :5]
    rng = np.random.default_rng(SEED)
    obs_mean_k5 = foreign_shares[5].mean()
    perm_means = np.empty(B_PERM)
    for b in range(B_PERM):
        perm_code = rng.permutation(code_arr)
        neighbor_labels = perm_code[nn5_idx]
        own_labels = perm_code[:, None]
        perm_means[b] = (neighbor_labels != own_labels).mean()
    perm_p = (1.0 + np.sum(perm_means <= obs_mean_k5)) / (B_PERM + 1)

    perm_info = {
        "obs_mean_foreign_share_k5": obs_mean_k5,
        "perm_mean_foreign_share_k5": perm_means.mean(),
        "perm_p_foreign_share_k5": perm_p,
        "perm_B": B_PERM,
    }

    return region_out, perm_info


# =============================================================================
# 4/7 — ANALYSIS 3: COUNTRY-AS-UNIT COUNTERFACTUAL
# =============================================================================

def run_counterfactual(
    t: dict[str, pd.DataFrame], region_neighbor: pd.DataFrame, region_order: pd.Series, code: pd.Series,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    fact_region = t["Fact_Region"].set_index("region_id").reindex(region_order)
    ccr = t["Country_Ref_Region"].set_index("region_id").reindex(region_order)

    hard_cluster = fact_region["hard_cluster"].astype(int)
    code_arr = code.to_numpy()

    majority_by_country: dict[str, int | None] = {}
    tie_countries: list[str] = []
    for c in np.unique(code_arr):
        counts = hard_cluster[code_arr == c].value_counts()
        top = counts[counts == counts.max()].index.tolist()
        if len(top) > 1:
            majority_by_country[c] = None
            tie_countries.append(c)
        else:
            majority_by_country[c] = int(top[0])

    country_majority_profile = pd.Series(
        [majority_by_country[c] for c in code_arr], index=region_order
    )
    differs_from_country_majority = (hard_cluster != country_majority_profile) & country_majority_profile.notna()

    region_extra = pd.DataFrame({
        "country_majority_profile": country_majority_profile,
        "differs_from_country_majority": differs_from_country_majority,
    })

    converged = fact_region["converged_toward_destination"].astype(bool)
    away = fact_region["moved_away_from_destination"].astype(bool)
    own_nearest = ccr["nearest_reference_is_own_country"].astype(bool)

    country_rows = []
    for c in np.unique(code_arr):
        mask = code_arr == c
        n_regions = int(mask.sum())
        profiles_present = int(hard_cluster[mask].nunique())
        minority_regions = int(differs_from_country_majority[mask].sum())
        regions_toward = int(converged[mask].sum())
        regions_away = int(away[mask].sum())
        has_both_directions = bool(regions_toward > 0 and regions_away > 0)

        neighbor_mask = region_neighbor["code"].to_numpy() == c
        mean_foreign_share_k1 = float(region_neighbor.loc[neighbor_mask, "foreign_share_k1"].mean())
        mean_foreign_share_k5 = float(region_neighbor.loc[neighbor_mask, "foreign_share_k5"].mean())
        mean_baseline_foreign_share = float(region_neighbor.loc[neighbor_mask, "baseline_foreign_share"].mean())

        country_rows.append({
            "code": c,
            "n_regions": n_regions,
            "profiles_present": profiles_present,
            "majority_profile": majority_by_country[c],
            "minority_regions": minority_regions,
            "regions_toward": regions_toward,
            "regions_away": regions_away,
            "has_both_directions": has_both_directions,
            "mean_foreign_share_k1": mean_foreign_share_k1,
            "mean_foreign_share_k5": mean_foreign_share_k5,
            "mean_baseline_foreign_share": mean_baseline_foreign_share,
            "own_country_nearest_share": float(own_nearest[mask].mean()),
        })

    country_out = pd.DataFrame(country_rows)

    if tie_countries:
        print(f"NOTE: hard_cluster mode tie detected for countries: {tie_countries} "
              "(majority_profile left null for their regions, per spec A.4).")

    return region_extra, country_out


# =============================================================================
# 5/7 — QA AGAINST SECTION A.6
# =============================================================================

def validate_against_a6(
    variance: pd.DataFrame, region_out: pd.DataFrame, country_out: pd.DataFrame,
    perm_info: dict[str, Any], distance_check: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    v = variance.set_index("variable")
    for name, (exp_eta2, exp_eps2) in EXPECTED_VARIANCE.items():
        obs_eta2 = float(v.loc[name, "eta2"])
        obs_eps2 = float(v.loc[name, "epsilon2"])
        qa_add(rows, f"eta2::{name}", abs(obs_eta2 - exp_eta2) <= QA_TOLERANCE,
               round(obs_eta2, 6), exp_eta2)
        qa_add(rows, f"epsilon2::{name}", abs(obs_eps2 - exp_eps2) <= QA_TOLERANCE,
               round(obs_eps2, 6), exp_eps2)

    null_eta2 = float(variance["null_expected_eta2"].iloc[0])
    qa_add(rows, "null_expected_eta2", abs(null_eta2 - EXPECTED_NULL_ETA2) <= QA_TOLERANCE,
           round(null_eta2, 6), EXPECTED_NULL_ETA2)

    max_diff = float(distance_check["abs_diff"].max())
    qa_add(rows, "distance_reproduces_region_country_ref",
           max_diff < NUMERIC_TOL_DISTANCE, max_diff, f"< {NUMERIC_TOL_DISTANCE}")
    qa_add(rows, "distance_reproduces_region_country_ref_rowcount",
           len(distance_check) == EXPECTED_REGIONS * EXPECTED_COUNTRIES,
           len(distance_check), EXPECTED_REGIONS * EXPECTED_COUNTRIES)

    nn1_foreign_count = int(region_out["nn1_is_foreign"].sum())
    nn1_foreign_share = nn1_foreign_count / len(region_out)
    qa_add(rows, "nn1_is_foreign_count", nn1_foreign_count == EXPECTED_NN1_FOREIGN_COUNT,
           nn1_foreign_count, EXPECTED_NN1_FOREIGN_COUNT)
    qa_add(rows, "nn1_is_foreign_share", abs(nn1_foreign_share - EXPECTED_NN1_FOREIGN_SHARE) <= QA_TOLERANCE,
           round(nn1_foreign_share, 3), EXPECTED_NN1_FOREIGN_SHARE)

    mean_k5 = float(region_out["foreign_share_k5"].mean())
    majority_k5 = int((region_out["foreign_share_k5"] > 0.5).sum())
    qa_add(rows, "mean_foreign_share_k5", abs(mean_k5 - EXPECTED_MEAN_FOREIGN_K5) <= QA_TOLERANCE,
           round(mean_k5, 3), EXPECTED_MEAN_FOREIGN_K5)
    qa_add(rows, "majority_foreign_k5_count", majority_k5 == EXPECTED_MAJORITY_FOREIGN_K5,
           majority_k5, EXPECTED_MAJORITY_FOREIGN_K5)

    mean_k10 = float(region_out["foreign_share_k10"].mean())
    majority_k10 = int((region_out["foreign_share_k10"] > 0.5).sum())
    qa_add(rows, "mean_foreign_share_k10", abs(mean_k10 - EXPECTED_MEAN_FOREIGN_K10) <= QA_TOLERANCE,
           round(mean_k10, 3), EXPECTED_MEAN_FOREIGN_K10)
    qa_add(rows, "majority_foreign_k10_count", majority_k10 == EXPECTED_MAJORITY_FOREIGN_K10,
           majority_k10, EXPECTED_MAJORITY_FOREIGN_K10)

    mean_baseline = float(region_out["baseline_foreign_share"].mean())
    qa_add(rows, "mean_baseline_foreign_share", abs(mean_baseline - EXPECTED_MEAN_BASELINE_FOREIGN) <= QA_TOLERANCE,
           round(mean_baseline, 3), EXPECTED_MEAN_BASELINE_FOREIGN)

    qa_add(rows, "perm_mean_foreign_share_k5_approx",
           abs(perm_info["perm_mean_foreign_share_k5"] - EXPECTED_PERM_MEAN_APPROX) <= 0.01,
           round(perm_info["perm_mean_foreign_share_k5"], 3), EXPECTED_PERM_MEAN_APPROX)
    qa_add(rows, "perm_p_foreign_share_k5_significant",
           perm_info["perm_p_foreign_share_k5"] <= EXPECTED_PERM_P_MAX,
           perm_info["perm_p_foreign_share_k5"], f"<= {EXPECTED_PERM_P_MAX}")

    differs_count = int(country_out["minority_regions"].sum())
    differs_pct = differs_count / EXPECTED_REGIONS
    qa_add(rows, "differs_from_country_majority_count", differs_count == EXPECTED_DIFFERS_MAJORITY,
           differs_count, EXPECTED_DIFFERS_MAJORITY)
    qa_add(rows, "differs_from_country_majority_pct", abs(differs_pct - EXPECTED_DIFFERS_MAJORITY_PCT) <= QA_TOLERANCE,
           round(differs_pct, 3), EXPECTED_DIFFERS_MAJORITY_PCT)

    both_directions = int(country_out["has_both_directions"].sum())
    qa_add(rows, "countries_with_both_directions", both_directions == EXPECTED_BOTH_DIRECTIONS_COUNTRIES,
           both_directions, EXPECTED_BOTH_DIRECTIONS_COUNTRIES)

    own_nearest_cross_check = int(round(country_out["own_country_nearest_share"].to_numpy() @
                                         country_out["n_regions"].to_numpy()))
    qa_add(rows, "own_country_nearest_cross_check", own_nearest_cross_check == EXPECTED_OWN_COUNTRY_NEAREST,
           own_nearest_cross_check, EXPECTED_OWN_COUNTRY_NEAREST)

    qa_add(rows, "rvc_region_rows", len(region_out) == EXPECTED_REGIONS, len(region_out), EXPECTED_REGIONS)
    qa_add(rows, "rvc_country_rows", len(country_out) == EXPECTED_COUNTRIES, len(country_out), EXPECTED_COUNTRIES)
    qa_add(rows, "rvc_variance_rows", len(variance) == 14, len(variance), 14)

    return enforce_qa(rows, "Script 07.3 QA against specification Section A.6 (stop point)")


# =============================================================================
# 6/7 — SUMMARY + REPORT
# =============================================================================

def build_summary(variance: pd.DataFrame, region_out: pd.DataFrame, country_out: pd.DataFrame,
                   perm_info: dict[str, Any]) -> pd.DataFrame:
    v = variance.set_index("variable")
    rows = [
        ("country_share_of_profile_variation_eta2", float(v.loc["hard_cluster_is_C1", "eta2"])),
        ("country_share_of_movement_variation_eta2", float(v.loc["convergence_index", "eta2"])),
        ("null_expected_eta2", float(variance["null_expected_eta2"].iloc[0])),
        ("regions_nearest_abroad_count", int(region_out["nn1_is_foreign"].sum())),
        ("regions_nearest_abroad_share", float(region_out["nn1_is_foreign"].mean())),
        ("regions_nearest_own_country_count", EXPECTED_REGIONS - int(region_out["nn1_is_foreign"].sum())),
        ("mean_foreign_share_k5", float(region_out["foreign_share_k5"].mean())),
        ("mean_baseline_foreign_share", float(region_out["baseline_foreign_share"].mean())),
        ("perm_p_foreign_share_k5", float(perm_info["perm_p_foreign_share_k5"])),
        ("regions_differ_from_country_majority", int(country_out["minority_regions"].sum())),
        ("countries_with_both_directions", int(country_out["has_both_directions"].sum())),
        ("n_countries", EXPECTED_COUNTRIES),
    ]
    return pd.DataFrame(rows, columns=["parameter", "value"])


def build_report(variance: pd.DataFrame, region_out: pd.DataFrame, country_out: pd.DataFrame,
                  summary: pd.DataFrame, perm_info: dict[str, Any], input_manifest: pd.DataFrame,
                  elapsed: float) -> str:
    s = summary.set_index("parameter")["value"]
    v = variance.set_index("variable")

    lines: list[str] = []
    lines.append("# Script 07.3 — Region vs country: results report")
    lines.append("")
    lines.append(f"Version: {SCRIPT_VERSION} | Generated: {pd.Timestamp.now().isoformat(timespec='seconds')}")
    lines.append("")
    lines.append(
        "Research question: does it pay off to compare regions rather than countries? "
        "This script decomposes 14 region-level variables into a between-country and a "
        "within-country share, measures how often a region's most similar region (by "
        "trajectory distance) is located abroad, and computes a country-as-unit "
        "counterfactual against the model's frozen hard clusters."
    )
    lines.append("")
    lines.append("## Headline findings")
    lines.append("")
    lines.append(
        f"- Country explains **{s['country_share_of_profile_variation_eta2']:.1%}** of the "
        "variation in the assigned profile (eta2 of hard_cluster_is_C1), against a chance "
        f"level of **{s['null_expected_eta2']:.1%}**."
    )
    lines.append(
        f"- Country explains only **{s['country_share_of_movement_variation_eta2']:.1%}** of "
        "the variation in the convergence index — movement is mostly within-country."
    )
    lines.append(
        f"- For **{int(s['regions_nearest_abroad_count'])} regions ({s['regions_nearest_abroad_share']:.1%})**, "
        "the single most similar region (by trajectory distance) is located in another country."
    )
    lines.append(
        f"- Across each region's 5 nearest neighbors, the average foreign share is "
        f"**{s['mean_foreign_share_k5']:.1%}**, against a structural baseline of "
        f"**{s['mean_baseline_foreign_share']:.1%}** if country did not matter at all "
        f"(permutation p ≤ {EXPECTED_PERM_P_MAX})."
    )
    lines.append(
        f"- Counterfactual: if every region inherited its country's majority profile, "
        f"**{int(s['regions_differ_from_country_majority'])} regions "
        f"({int(s['regions_differ_from_country_majority'])/EXPECTED_REGIONS:.1%})** would end up "
        "with a different profile than their own."
    )
    lines.append(
        f"- **{int(s['countries_with_both_directions'])} of {EXPECTED_COUNTRIES} countries** "
        "contain regions moving toward the alternative profile AND regions moving away from "
        "it at the same time."
    )
    lines.append("")
    lines.append(
        "Suggested formulation (one interpretation, not a result): the global partition "
        "distinguishes mostly deprivation levels that vary between countries; the subnational "
        "dimension shows up mainly in trajectories, membership gradation, and similarity "
        "between regions of different countries."
    )
    lines.append("")
    lines.append("## RVC_Variance — how much does the country explain (Section A.2)")
    lines.append("")
    lines.append("| Group | Variable | eta2 | epsilon2 | within_share | perm_p |")
    lines.append("|---|---|---:|---:|---:|---:|")
    for _, r in variance.iterrows():
        lines.append(
            f"| {r['variable_group']} | {r['variable_label']} | {r['eta2']:.3f} | "
            f"{r['epsilon2']:.3f} | {r['within_share']:.3f} | {r['perm_p']:.4f} |"
        )
    lines.append("")
    lines.append(f"Chance-level eta2 ((k-1)/(N-1)): **{v['null_expected_eta2'].iloc[0] if 'null_expected_eta2' in v.columns else s['null_expected_eta2']:.4f}**.")
    lines.append("")
    lines.append("## RVC_Region / RVC_Country — neighbors across borders (Section A.3)")
    lines.append("")
    lines.append(f"- Nearest region abroad (k=1): {int(s['regions_nearest_abroad_count'])} regions "
                  f"({s['regions_nearest_abroad_share']:.1%}); nearest region in own country: "
                  f"{int(s['regions_nearest_own_country_count'])} regions.")
    lines.append(f"- Mean foreign share among 5 nearest neighbors: {s['mean_foreign_share_k5']:.1%} "
                  f"(structural baseline {s['mean_baseline_foreign_share']:.1%}).")
    lines.append(
        f"- Permutation test (B={perm_info['perm_B']:,}): permutation mean "
        f"{perm_info['perm_mean_foreign_share_k5']:.3f} vs observed {perm_info['obs_mean_foreign_share_k5']:.3f}, "
        f"p = {perm_info['perm_p_foreign_share_k5']:.4f}."
    )
    lines.append("")
    lines.append("## RVC_Country — country-as-unit counterfactual (Section A.4)")
    lines.append("")
    lines.append(
        f"- {int(s['regions_differ_from_country_majority'])} of {EXPECTED_REGIONS} regions "
        f"({int(s['regions_differ_from_country_majority'])/EXPECTED_REGIONS:.1%}) would end up with a "
        "different profile if their country's majority profile were used instead of their own."
    )
    lines.append(
        f"- {int(s['countries_with_both_directions'])} of {EXPECTED_COUNTRIES} countries contain regions "
        "moving toward the alternative profile and regions moving away from it at the same time."
    )
    lines.append("")
    lines.append("## Caveats (must be preserved in any report/dashboard text built on this)")
    lines.append("")
    lines.append(
        "- eta2 is descriptive and does not weight by fuzzy membership. Countries with 3-4 "
        "regions inflate the between-country share (hence epsilon2 as a bias-corrected companion)."
    )
    lines.append("- Units are mixed (GAUL and NUTS).")
    lines.append("- Trajectory proximity does not imply geographic proximity nor causality.")
    lines.append(
        "- Never phrase results as \"regions are better than countries\" or \"countries are "
        "irrelevant\"; never use causal language (\"because of\", \"driven by\"); never treat "
        "these results as proof that the global model is wrong; always mention the structural "
        "minimum of small countries when discussing foreign share; never compare a foreign "
        "share without its baseline."
    )
    lines.append("")
    lines.append("## Inputs")
    lines.append("")
    for _, r in input_manifest.iterrows():
        lines.append(f"- `{r['relative_path']}` — sha256 `{r['sha256']}`")
    lines.append("")
    lines.append(f"Script execution time: {elapsed:.2f} seconds.")
    lines.append("")
    lines.append(
        "Pending decision (per spec Section A.5): whether to fold Script 07.3 into the "
        "07.1/07.2 consolidation is left to the author; scripts 07.1 and 07.2 were not modified."
    )
    lines.append("")
    return "\n".join(lines)


# =============================================================================
# 7/7 — WRITE OUTPUTS
# =============================================================================

def build_metadata(input_manifest: pd.DataFrame, elapsed: float) -> pd.DataFrame:
    rows = [
        {"key": "script", "value": "07.3_region_vs_country.py"},
        {"key": "script_version", "value": SCRIPT_VERSION},
        {"key": "generated_at", "value": pd.Timestamp.now().isoformat(timespec="seconds")},
        {"key": "seed", "value": SEED},
        {"key": "perm_B", "value": B_PERM},
        {"key": "k_list", "value": str(K_LIST)},
        {"key": "n_regions", "value": EXPECTED_REGIONS},
        {"key": "n_countries", "value": EXPECTED_COUNTRIES},
        {"key": "input_workbook", "value": input_manifest["relative_path"].iloc[0]},
        {"key": "input_workbook_sha256", "value": input_manifest["sha256"].iloc[0]},
        {"key": "execution_seconds", "value": round(elapsed, 2)},
        {"key": "note", "value": "No re-estimation of the frozen analytical model. Scripts 01-07.2 and outputs/ untouched."},
    ]
    return pd.DataFrame(rows)


def write_workbook(variance: pd.DataFrame, region_out: pd.DataFrame, country_out: pd.DataFrame,
                    summary: pd.DataFrame, metadata: pd.DataFrame) -> None:
    with pd.ExcelWriter(WORKBOOK_OUTPUT, engine="openpyxl") as writer:
        variance.to_excel(writer, sheet_name="RVC_Variance", index=False)
        region_out.to_excel(writer, sheet_name="RVC_Region", index=False)
        country_out.to_excel(writer, sheet_name="RVC_Country", index=False)
        summary.to_excel(writer, sheet_name="RVC_Summary", index=False)
        metadata.to_excel(writer, sheet_name="RVC_Metadata", index=False)
    autofit_workbook(WORKBOOK_OUTPUT)


# =============================================================================
# MAIN
# =============================================================================

def main() -> None:
    t0 = time.perf_counter()
    ensure_directories()

    print_header("SCRIPT 07.3 v1.0.0 — REGION VS COUNTRY")
    print(f"Project root: {ROOT}")
    print("Analytical recalculation: NONE (reads only outputs/07.2_powerbi_model.xlsx)")
    print(f"Seed: {SEED} | Permutations: {B_PERM:,} | k_list: {K_LIST}")

    print_header("1/7 — VALIDATE INPUT CONTRACT")
    input_manifest = build_input_manifest()
    t = load_inputs()
    validate_input_contract(t)
    print(f"Source workbook: {relative_path(SOURCE_WORKBOOK)}")
    print(f"Source sha256:   {input_manifest['sha256'].iloc[0]}")

    region_order = t["Fact_Region"].sort_values("region_id", kind="stable")["region_id"].reset_index(drop=True)
    code = (
        t["Fact_Region"].set_index("region_id")["code"].reindex(region_order).reset_index(drop=True)
    )
    print(f"Canonical region order: {len(region_order)} regions, {code.nunique()} countries.")

    print_header("2/7 — ANALYSIS 1: VARIANCE DECOMPOSITION")
    matrix = build_variable_matrix(t, region_order)
    variance = run_variance_decomposition(matrix, code)
    print(variance[["variable", "eta2", "epsilon2", "perm_p"]].to_string(index=False))

    print_header("3/7 — ANALYSIS 2: NEIGHBORS ACROSS BORDERS")
    F = build_feature_year_matrix(t, region_order)
    distance_check = validate_country_reference_distance(F, region_order, code, t)
    print(f"Max |diff| vs Region_Country_Ref.mean_annual_distance_2015_2019: "
          f"{distance_check['abs_diff'].max():.3e} (tolerance < {NUMERIC_TOL_DISTANCE})")
    region_neighbor, perm_info = run_neighbor_analysis(F, region_order, code, t)
    print(f"nn1_is_foreign: {int(region_neighbor['nn1_is_foreign'].sum())} regions")
    print(f"mean foreign_share_k5: {region_neighbor['foreign_share_k5'].mean():.4f} "
          f"(baseline {region_neighbor['baseline_foreign_share'].mean():.4f})")
    print(f"permutation p (foreign_share_k5): {perm_info['perm_p_foreign_share_k5']:.4f}")

    print_header("4/7 — ANALYSIS 3: COUNTRY-AS-UNIT COUNTERFACTUAL")
    region_extra, country_out = run_counterfactual(t, region_neighbor, region_order, code)
    region_out = pd.concat([region_neighbor.set_index("region_id"), region_extra], axis=1).reset_index()
    print(f"differs_from_country_majority: {int(region_extra['differs_from_country_majority'].sum())} regions")
    print(f"countries with both directions: {int(country_out['has_both_directions'].sum())} of {EXPECTED_COUNTRIES}")

    print_header("5/7 — QA AGAINST SPECIFICATION SECTION A.6 (STOP POINT)")
    qa_frame = validate_against_a6(variance, region_out, country_out, perm_info, distance_check)
    print(f"All {len(qa_frame)} QA checks against Section A.6: PASS")

    print_header("6/7 — BUILD SUMMARY + REPORT")
    summary = build_summary(variance, region_out, country_out, perm_info)
    elapsed_preview = time.perf_counter() - t0
    metadata = build_metadata(input_manifest, elapsed_preview)
    report_text = build_report(variance, region_out, country_out, summary, perm_info, input_manifest, elapsed_preview)

    print_header("7/7 — WRITE OUTPUTS")
    write_workbook(variance, region_out, country_out, summary, metadata)
    REPORT_OUTPUT.write_text(report_text, encoding="utf-8")
    write_csv(qa_frame, ROOT / "data" / "interim" / "script07" / "07.3_qa.csv")

    status = {
        "script": "07.3_region_vs_country.py",
        "script_version": SCRIPT_VERSION,
        "status": "PASS",
        "qa_checks": len(qa_frame),
        "regions": EXPECTED_REGIONS,
        "countries": EXPECTED_COUNTRIES,
        "seed": SEED,
        "perm_B": B_PERM,
        "outputs": {
            "workbook": relative_path(WORKBOOK_OUTPUT),
            "report": relative_path(REPORT_OUTPUT),
        },
        "prior_analytical_recalculation": False,
        "scripts_01_to_07_2_modified": False,
        "outputs_dir_pre_existing_files_modified": False,
        "part_b_pending": True,
    }
    FINAL_STATUS_OUTPUT.write_text(json.dumps(status, indent=2, ensure_ascii=False), encoding="utf-8")

    elapsed = time.perf_counter() - t0
    print(f"Workbook: {relative_path(WORKBOOK_OUTPUT)}")
    print(f"Report:   {relative_path(REPORT_OUTPUT)}")
    print(f"Status:   {relative_path(FINAL_STATUS_OUTPUT)}")

    print("\nFINAL SCRIPT 07.3 SUMMARY")
    print("-" * 100)
    print(f"Regions:                {EXPECTED_REGIONS}")
    print(f"Countries:              {EXPECTED_COUNTRIES}")
    print(f"QA checks vs Section A.6: {len(qa_frame)} PASS")
    print("Prior analytical recalculation: NONE")
    print("\nSCRIPT 07.3 COMPLETE — PART A QA PASSED; PART B (dashboard copy) MAY PROCEED")
    print(f"Script execution time: {elapsed:.2f} seconds")


if __name__ == "__main__":
    main()
